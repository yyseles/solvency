# -*- coding: utf-8 -*-
"""用「完整数据源」重建 data.js 再保(reins)板块全部公司数据。

数据源（用户提供，已含中再集团本部）:
  data/再保偿付能力22Q1-26Q2.xlsx    -> data 主字段 C..V（16家 x 22期：季度18 + 年度4）
  data/再保最低资本22Q1-26Q2.xlsx    -> mcDetail（16家 x 季度期次，切割到 2026Q2，去 2026Q3/Q4 空期次）

铁律：
  1) 只动 reins 段，绝不碰 group/property/life；
  2) 公司名用数据源名「中再集团本部」（替换旧名「中再集团本级」，不保留）；
  3) mcDetail 只保留存在于 data 的期次（<=2026Q2），丢弃空期次；
  4) 保留 data.js 首行注释 + let 前缀。
"""
import os, re, sys
import openpyxl
import json5

HERE = os.path.dirname(os.path.abspath(__file__))
SOLV = os.path.dirname(HERE)
DATA_JS = os.path.join(SOLV, "data.js")
DATA_DIR = os.path.join(SOLV, "data")
SOLV_FILE = os.path.join(DATA_DIR, "再保偿付能力22Q1-26Q2.xlsx")
MC_FILE = os.path.join(DATA_DIR, "再保最低资本22Q1-26Q2.xlsx")

CN = {"一": "1", "二": "2", "三": "3", "四": "4"}


def normalize_period(s):
    s = str(s).strip()
    m = re.match(r"(\d{4})第?([一二三四1-4])季度", s)
    if m:
        q = CN.get(m.group(2), m.group(2))
        return "%sQ%s" % (m.group(1), q)
    m = re.match(r"(\d{4})", s)
    return m.group(1)


# 偿付能力文件：列标题(row2) -> 字段
SOLV_COL_MAP = {
    "综合偿付能力充足率": "C", "核心偿付能力充足率": "D",
    "综合偿付能力溢额": "E", "核心偿付能力溢额": "F",
    "认可资产": "G", "认可负债": "H", "实际资本": "I",
    "核心一级资本": "J", "核心二级资本": "K", "附属一级资本": "L", "附属二级资本": "M",
    "最低资本": "N", "量化风险最低资本": "O", "寿险风险最低资本": "P",
    "非寿险风险最低资本": "Q", "市场风险最低资本": "R", "信用风险最低资本": "S",
    "量化风险分散效应": "T", "特定类别保险合同损失吸收效应": "U", "控制风险最低资本": "V",
}

# 最低资本文件：行标签 -> mc 字段
MC_METRIC_MAP = {
    "寿险业务保险风险-损失发生风险最低资本": "mcP_loss",
    "寿险业务保险风险-退保风险最低资本": "mcP_surr",
    "寿险业务保险风险-费用风险最低资本": "mcP_exp",
    "寿险业务保险风险-风险分散效应": "mcP_div",
    "非寿险业务保险风险-保费及准备金风险最低资本": "mcQ_prem",
    "非寿险业务保险风险-巨灾风险最低资本": "mcQ_cata",
    "非寿险业务保险风险-风险分散效应": "mcQ_div",
    "市场风险-利率风险最低资本": "mcR_rate",
    "市场风险-权益价格风险最低资本": "mcR_eq",
    "市场风险-房地产价格风险最低资本": "mcR_re",
    "市场风险-境外固定收益类资产价格风险最低资本": "mcR_ofb",
    "市场风险-境外权益类资产价格风险最低资本": "mcR_ofe",
    "市场风险-汇率风险最低资本": "mcR_fx",
    "市场风险-风险分散效应": "mcR_div",
    "信用风险-利差风险最低资本": "mcS_spr",
    "信用风险-交易对手违约风险最低资本": "mcS_def",
    "信用风险-风险分散效应": "mcS_div",
    "附加资本": "mcAdd",
    "逆周期附加资本": "mcAdd_cyc",
    "D-SII附加资本": "mcAdd_dsii",
    "G-SII附加资本": "mcAdd_gsii",
    "其他附加资本": "mcAdd_oth",
}


def read_solvency(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["风险数据"]
    # 列索引
    col_idx = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=2, column=c).value
        if v and str(v).strip() in SOLV_COL_MAP:
            col_idx[SOLV_COL_MAP[str(v).strip()]] = c
    data = {}
    for r in range(3, ws.max_row + 1):
        c1 = ws.cell(row=r, column=1).value
        c2 = ws.cell(row=r, column=2).value
        if not c1 or not c2:
            continue
        comp = str(c1).strip()
        pk = normalize_period(c2)
        vals = {}
        for f, ci in col_idx.items():
            v = ws.cell(row=r, column=ci).value
            if v is not None:
                try:
                    vals[f] = float(v)
                except (TypeError, ValueError):
                    pass
        if vals:
            data.setdefault(comp, {})[pk] = vals
    wb.close()
    return data


def read_mincap(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    # 列->公司（row0 块起点延续）
    col_company = {}
    cur = None
    for ci in range(ws.max_column):
        c = ws.cell(row=1, column=ci + 1).value
        if c and str(c).strip():
            cur = str(c).strip()
        col_company[ci] = cur
    # 列->期次（row1）
    col_period = {}
    for ci in range(ws.max_column):
        p = ws.cell(row=2, column=ci + 1).value
        if p:
            col_period[ci] = normalize_period(str(p).strip())
    data = {}
    for ri in range(3, ws.max_row + 1):
        lab = ws.cell(row=ri, column=1).value
        if not lab:
            continue
        field = MC_METRIC_MAP.get(str(lab).strip())
        if not field:
            continue
        for ci in range(1, ws.max_column):
            comp = col_company.get(ci)
            pk = col_period.get(ci)
            if not comp or not pk:
                continue
            v = ws.cell(row=ri, column=ci + 1).value
            if v is None:
                continue
            try:
                val = float(v)
            except (TypeError, ValueError):
                continue
            data.setdefault(comp, {}).setdefault(pk, {})[field] = val
    wb.close()
    return data


def make_period_def(key):
    m = re.match(r"(\d{4})Q([1-4])$", key)
    if m:
        y, q = int(m.group(1)), int(m.group(2))
        return {"key": key, "year": y, "q": q, "kind": "quarter",
                "label": "%s年第%s季度" % (y, q), "source": "quarter"}
    m = re.match(r"(\d{4})$", key)
    if m:
        y = int(m.group(1))
        return {"key": key, "year": y, "q": 4, "kind": "year-end",
                "label": "%s年" % y, "source": "audited"}
    return {"key": key, "year": 0, "q": 0, "kind": "other", "label": key, "source": "other"}


def main():
    solv = read_solvency(SOLV_FILE)
    mc = read_mincap(MC_FILE)
    print("偿付能力: %d 家公司" % len(solv))
    print("最低资本: %d 家公司" % len(mc))

    # data.js 现有结构
    txt = open(DATA_JS, encoding="utf-8").read()
    m = re.search(r"(?:let|const)\s+(\w+)\s*=\s*(\{.*\})\s*;", txt, re.S)
    var, D, header = m.group(1), json5.loads(m.group(2)), txt.split("\n", 1)[0]

    seg = D["segments"]["reins"]
    old_companies = seg["companies"]

    # 1) data 主字段，来自偿付能力文件
    new_data = dict(solv)  # 每公司 {pk: {C..V}}
    # 1b) 年度主数据兜底：源 Excel 年度行=Q4 值（其他15家已如此）；本部年度行为空，
    #     同口径用 Q4 补齐，保证 16家×22期 完整
    for comp, pd in new_data.items():
        for y in ["2022", "2023", "2024", "2025"]:
            if y not in pd and (y + "Q4") in pd:
                pd[y] = dict(pd[y + "Q4"])
    # 2) mcDetail：只保留存在于 data 的期次（<=2026Q2），丢弃空期次
    #    每公司 mcDetail 期次 = 该公司 data 期次 与 mc 期次 的交集
    new_mcd = {}
    for comp, mp in mc.items():
        dkeys = set(new_data.get(comp, {}).keys())
        kept = {}
        for pk, vals in mp.items():
            if pk in dkeys:  # 只保留 data 中有该期次的（含年度不含时也按交集）
                kept[pk] = vals
        if kept:
            new_mcd[comp] = kept

    # 2b) 年度 mcDetail：最低资本 Excel 只有季度明细；年度下钻明细按用户口径用 Q4 值
    #     （年度聚合值 O/P/Q..V/N 已由偿付能力 Excel 年度行提供，此处仅补下钻子项）
    for comp, pd in new_data.items():
        for pk in pd:
            if re.match(r"^\d{4}$", pk):
                q4 = new_mcd.get(comp, {}).get(pk + "Q4")
                if q4:
                    new_mcd.setdefault(comp, {})[pk] = dict(q4)

    # 3) companies 清单 = 偿付能力文件公司清单（数据源名）
    new_companies = list(new_data.keys())

    # 4) periods：从新 data 所有期次构建；尽力保留原 source/kind
    all_keys = set()
    for comp, pd in new_data.items():
        for pk in pd:
            all_keys.add(pk)
    old_defs = {p["key"]: p for p in seg["periods"]}
    new_periods = []
    for k in sorted(all_keys, key=lambda x: (int(x[:4]) if x[:4].isdigit() else 0, x)):
        if k in old_defs:
            new_periods.append(old_defs[k])
        else:
            new_periods.append(make_period_def(k))
    new_periods.sort(key=lambda x: (x.get("year", 0), x.get("q", 0)))

    # 5) 覆盖 rein 段
    seg["companies"] = new_companies
    seg["data"] = new_data
    seg["mcDetail"] = new_mcd
    seg["periods"] = new_periods

    # 写回
    js = header + "\nlet " + var + " = " + json5.dumps(
        D, ensure_ascii=False, separators=(",", ":")) + ";\n"
    tmp = DATA_JS + ".tmp"
    open(tmp, "w", encoding="utf-8").write(js)
    os.replace(tmp, DATA_JS)

    print("\n=== 重建完成 ===")
    print("公司清单(%d):" % len(new_companies), new_companies)
    print("期次数: %d" % len(new_periods),
          [p["key"] for p in new_periods])


if __name__ == "__main__":
    main()
